import zipfile

from click.testing import CliRunner
from pptx import Presentation

from atulya_office import core
from atulya_office.cli import main


def _make_docx(path, body):
    xml = ('<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/'
           'wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>%s</w:t></w:r></w:p>'
           '</w:body></w:document>' % body)
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("word/document.xml", xml)


def test_docx_placeholder_replace_and_escape(tmp_path):
    src, out = tmp_path / "t.docx", tmp_path / "o.docx"
    _make_docx(src, "Hi {{name}}")
    core._replace_docx_placeholders(str(src), {"name": "A & B"}, str(out))
    assert core._get_docx_text(str(out)) == "Hi A & B"


def test_build_ppt_from_outline(tmp_path):
    ol = tmp_path / "o.txt"
    ol.write_text("# One\n- a\n  - b\n# Two\n* c\n")
    out = core.build_ppt_from_outline(str(ol), str(tmp_path / "d.pptx"))
    prs = Presentation(out)
    assert len(prs.slides) == 2
    paras = prs.slides[0].placeholders[1].text_frame.paragraphs
    assert [p.level for p in paras] == [0, 1]


def test_cli_ppt_build_rejects_empty(tmp_path):
    ol = tmp_path / "o.txt"
    ol.write_text("no slides here\n")
    r = CliRunner().invoke(main, ["ppt", "build", str(ol), "-o", str(tmp_path / "x.pptx")])
    assert r.exit_code != 0
